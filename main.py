import os
import io
import json
import logging
import asyncio
import sqlite3
from datetime import date, datetime
from typing import List, Optional

# Async & Document Libraries
import aiofiles
from pypdf import PdfReader
from docx import Document
from pptx import Presentation
import openpyxl

# AI & Pydantic
from pydantic import BaseModel, Field
from google import genai
from google.genai import types
from tenacity import retry, stop_after_attempt, wait_exponential, retry_if_exception_type

# Telegram Bot Engine (aiogram 3.x)
from aiogram import Bot, Dispatcher, F, Router
from aiogram.filters import Command, CommandStart
from aiogram.types import (
    Message, CallbackQuery, InlineKeyboardMarkup, InlineKeyboardButton,
    WebAppInfo, ContentType
)
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.fsm.storage.memory import MemoryStorage

# Logging Setup
logging.basicConfig(level=logging.INFO, format="%(asctime)s - %(levelname)s - %(message)s")

# ==========================================
# CONFIGURATION & ENVIRONMENT VARIABLES
# ==========================================
TELEGRAM_TOKEN = os.getenv("TELEGRAM_TOKEN", "YOUR_TELEGRAM_BOT_TOKEN")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "YOUR_GEMINI_API_KEY")
INITIAL_ADMIN_ID = int(os.getenv("ADMIN_ID", "123456789"))

# GitHub Pages ဖြင့် Hosting တင်ထားသော သင်၏ WebApp Link
WEBAPP_URL = "https://cnbintell.github.io/telegram-quiz-bot/quiz_webapp.html"

# Initialize Gemini Client
ai_client = genai.Client(api_key=GEMINI_API_KEY)

# ==========================================
# 1. DATABASE MANAGEMENT (SQLite)
# ==========================================
DB_FILE = "enterprise_quiz.db"

def init_db():
    conn = sqlite3.connect(DB_FILE)
    cursor = conn.cursor()
    
    # User Table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS User (
        user_id INTEGER PRIMARY KEY,
        username TEXT,
        is_premium BOOLEAN DEFAULT 0,
        daily_count INTEGER DEFAULT 0,
        last_quiz_date TEXT,
        created_at TEXT
    )
    """)
    
    # Admin Table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS Admin (
        admin_id INTEGER PRIMARY KEY,
        added_by INTEGER,
        created_at TEXT
    )
    """)
    
    # QuizCategory Table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS QuizCategory (
        category_id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT UNIQUE,
        time_limit_seconds INTEGER DEFAULT 30
    )
    """)
    
    # PaymentConfig Table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS PaymentConfig (
        config_id INTEGER PRIMARY KEY DEFAULT 1,
        kpay_number TEXT,
        kpay_name TEXT,
        qr_code_file_id TEXT,
        updated_at TEXT
    )
    """)
    
    # QuizResult Table
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS QuizResult (
        result_id INTEGER PRIMARY KEY AUTOINCREMENT,
        user_id INTEGER,
        category_name TEXT,
        score INTEGER,
        total_questions INTEGER,
        time_taken_seconds INTEGER,
        submitted_at TEXT
    )
    """)

    # Questions Storage
    cursor.execute("""
    CREATE TABLE IF NOT EXISTS QuestionBank (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        category_name TEXT,
        question TEXT,
        options TEXT,
        correct_index INTEGER,
        explanation TEXT
    )
    """)
    
    # Default Initial Setup
    cursor.execute("INSERT OR IGNORE INTO Admin (admin_id, added_by, created_at) VALUES (?, ?, ?)",
                   (INITIAL_ADMIN_ID, 0, datetime.now().isoformat()))
    cursor.execute("INSERT OR IGNORE INTO QuizCategory (name, time_limit_seconds) VALUES ('General', 30)")
    cursor.execute("INSERT OR IGNORE INTO PaymentConfig (config_id, kpay_number, kpay_name, qr_code_file_id, updated_at) VALUES (1, '09123456789', 'Admin KPay', '', ?)",
                   (datetime.now().isoformat(),))
    
    conn.commit()
    conn.close()

init_db()

# DB Helper Async Operations
async def db_query(query: str, params: tuple = (), fetchone=False, fetchall=False, commit=False):
    def _execute():
        conn = sqlite3.connect(DB_FILE)
        cursor = conn.cursor()
        cursor.execute(query, params)
        res = None
        if fetchone:
            res = cursor.fetchone()
        elif fetchall:
            res = cursor.fetchall()
        if commit:
            conn.commit()
        conn.close()
        return res
    return await asyncio.to_thread(_execute)

# ==========================================
# 2. PYDANTIC SCHEMAS FOR STRUCTURED OUTPUT
# ==========================================
class QuizQuestionSchema(BaseModel):
    question: str = Field(description="The quiz question text")
    options: List[str] = Field(description="List of 4 multiple choice options")
    correct_index: int = Field(description="Zero-based index of the correct option (0-3)")
    explanation: str = Field(description="Explanation of why the correct answer is right")

class QuizSetSchema(BaseModel):
    questions: List[QuizQuestionSchema] = Field(description="List of extracted quiz questions")

# ==========================================
# 3. HIGH-RESILIENCE AI INGESTION SUBSYSTEM
# ==========================================
def extract_text_from_bytes(file_bytes: bytes, file_name: str) -> str:
    extracted_text = ""
    file_name = file_name.lower()
    
    if file_name.endswith('.pdf'):
        reader = PdfReader(io.BytesIO(file_bytes))
        for page in reader.pages:
            text = page.extract_text()
            if text:
                extracted_text += text + "\n"
    elif file_name.endswith('.docx'):
        doc = Document(io.BytesIO(file_bytes))
        extracted_text = "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
    elif file_name.endswith('.pptx'):
        prs = Presentation(io.BytesIO(file_bytes))
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    extracted_text += shape.text + "\n"
    elif file_name.endswith('.xlsx'):
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                extracted_text += " | ".join([str(c) for c in row if c is not None]) + "\n"
    elif file_name.endswith('.txt'):
        extracted_text = file_bytes.decode('utf-8', errors='ignore')
        
    return extracted_text

@retry(
    reraise=True,
    stop=stop_after_attempt(5),
    wait=wait_exponential(multiplier=1.0, min=2, max=32),
    retry=retry_if_exception_type(Exception)
)
async def generate_quiz_from_gemini(input_content: str, count: int = 5) -> QuizSetSchema:
    prompt = f"Extract or generate exactly {count} multiple choice questions based on the content below.\n\nContent:\n{input_content[:8000]}"
    
    response = await asyncio.to_thread(
        ai_client.models.generate_content,
        model='gemini-2.5-flash',
        contents=prompt,
        config=types.GenerateContentConfig(
            response_mime_type="application/json",
            response_schema=QuizSetSchema,
            temperature=0.2,
        ),
    )
    
    raw_text = response.text
    parsed_json = json.loads(raw_text)
    return QuizSetSchema(**parsed_json)

# ==========================================
# 4. TELEGRAM BOT ENGINE & HANDLERS
# ==========================================
router = Router()

class AdminStates(StatesGroup):
    wait_for_file_upload = State()

def get_main_menu(user_id: int, is_admin: bool):
    buttons = [
        [InlineKeyboardButton(text="🎯 Open WebApp Quiz", web_app=WebAppInfo(url=WEBAPP_URL))],
        [InlineKeyboardButton(text="💎 Upgrade Premium / Status", callback_data="check_status")]
    ]
    if is_admin:
        buttons.append([InlineKeyboardButton(text="👑 Admin Control Panel", callback_data="admin_panel")])
    return InlineKeyboardMarkup(inline_keyboard=buttons)

def get_admin_menu():
    return InlineKeyboardMarkup(inline_keyboard=[
        [InlineKeyboardButton(text="📄 File to Quiz (AI)", callback_data="admin_ai_file")],
        [InlineKeyboardButton(text="🔙 Back to Main Menu", callback_data="main_menu")]
    ])

@router.message(CommandStart())
async def start_cmd(message: Message):
    user_id = message.from_user.id
    username = message.from_user.username or "User"
    today = str(date.today())
    
    existing = await db_query("SELECT user_id FROM User WHERE user_id = ?", (user_id,), fetchone=True)
    if not existing:
        await db_query(
            "INSERT INTO User (user_id, username, is_premium, daily_count, last_quiz_date, created_at) VALUES (?, ?, 0, 0, ?, ?)",
            (user_id, username, today, datetime.now().isoformat()), commit=True
        )
    
    admin_row = await db_query("SELECT admin_id FROM Admin WHERE admin_id = ?", (user_id,), fetchone=True)
    is_admin = bool(admin_row)
    
    await message.answer(
        f"👋 **မင်္ဂလာပါ {username}**\nEnterprise-Grade AI Quiz Platform မှ ကြိုဆိုပါသည်။",
        reply_markup=get_main_menu(user_id, is_admin),
        parse_mode="Markdown"
    )

@router.callback_query(F.data == "main_menu")
async def cb_main_menu(callback: CallbackQuery):
    user_id = callback.from_user.id
    admin_row = await db_query("SELECT admin_id FROM Admin WHERE admin_id = ?", (user_id,), fetchone=True)
    await callback.message.edit_text("📌 **Main Menu**", reply_markup=get_main_menu(user_id, bool(admin_row)))

@router.callback_query(F.data == "admin_panel")
async def cb_admin_panel(callback: CallbackQuery):
    user_id = callback.from_user.id
    admin_row = await db_query("SELECT admin_id FROM Admin WHERE admin_id = ?", (user_id,), fetchone=True)
    if not admin_row:
        await callback.answer("❌ ခွင့်ပြုချက်မရှိပါ", show_alert=True)
        return
    await callback.message.edit_text("👑 **Admin Control Panel**", reply_markup=get_admin_menu())

@router.callback_query(F.data == "check_status")
async def cb_check_status(callback: CallbackQuery):
    user_id = callback.from_user.id
    user_row = await db_query("SELECT is_premium, daily_count, last_quiz_date FROM User WHERE user_id = ?", (user_id,), fetchone=True)
    pay_row = await db_query("SELECT kpay_number, kpay_name, qr_code_file_id FROM PaymentConfig WHERE config_id = 1", fetchone=True)
    
    is_prem = user_row[0] if user_row else 0
    status_str = "💎 **Premium Member** (Unlimited)" if is_prem else "🆓 **Free Member** (၁ ရက် ၁၀ ပုဒ်)"
    
    msg = f"👤 **အကောင့်အခြေအနေ**\n\nID: `{user_id}`\nအဆင့်: {status_str}\n\n"
    if not is_prem and pay_row:
        msg += f"✨ **Premium သို့ မြှင့်တင်ရန် ပေးချေရမည့် အချက်အလက်များ:**\n" \
               f"📱 KPay: `{pay_row[0]}` ({pay_row[1]})\n\n" \
               f"ငွေလွှဲပြီးပါက **Transaction Screenshot ပြေစာ** ကို ဤ Bot ထံ တိုက်ရိုက် ပေးပို့ပါ။"
            
    await callback.message.edit_text(msg, reply_markup=InlineKeyboardMarkup(inline_keyboard=[[InlineKeyboardButton(text="🔙 Back", callback_data="main_menu")]]), parse_mode="Markdown")

@router.message(F.photo)
async def handle_payment_screenshot(message: Message, bot: Bot):
    user_id = message.from_user.id
    photo_file_id = message.photo[-1].file_id
    
    await db_query("UPDATE User SET is_premium = 1 WHERE user_id = ?", (user_id,), commit=True)
    await message.reply("✅ **ငွေလွှဲပြေစာ လက်ခံရရှိပါသည်။**\nသင့်အကောင့်ကို Premium အဖြစ် အလိုအလျောက် မြှင့်တင်ပေးလိုက်ပါပြီ။")
    
    admins = await db_query("SELECT admin_id FROM Admin", fetchall=True)
    for admin in admins:
        try:
            revoke_kb = InlineKeyboardMarkup(inline_keyboard=[[
                InlineKeyboardButton(text="❌ Revoke Premium", callback_data=f"revoke_prem_{user_id}")
            ]])
            await bot.send_photo(
                chat_id=admin[0],
                photo=photo_file_id,
                caption=f"🔔 **ငွေလွှဲပြေစာ အသစ်ရောက်ရှိလာပါသည်။**\nUser ID: `{user_id}` (@{message.from_user.username or 'N/A'})\n\nစနစ်မှ Auto Premium ပေးထားပါသည်။ မှားယွင်းပါက Revoke နှိပ်ပါ။",
                reply_markup=revoke_kb,
                parse_mode="Markdown"
            )
        except Exception:
            pass

@router.callback_query(F.data.startswith("revoke_prem_"))
async def cb_revoke_premium(callback: CallbackQuery):
    target_user_id = int(callback.data.replace("revoke_prem_", ""))
    await db_query("UPDATE User SET is_premium = 0 WHERE user_id = ?", (target_user_id,), commit=True)
    await callback.answer("✅ Premium ပြန်လည် ရုပ်သိမ်းလိုက်ပါပြီ", show_alert=True)
    await callback.message.edit_caption(caption=callback.message.caption + "\n\n⛔ **[ADMIN ACTION: PREMIUM REVOKED]**")

@router.callback_query(F.data == "admin_ai_file")
async def cb_admin_ai_file(callback: CallbackQuery, state: FSMContext):
    await callback.message.edit_text("📄 **မေးခွန်းထုတ်လိုသော ဖိုင် (PDF, DOCX, PPTX, XLSX, TXT) ကို ပို့ပေးပါ -**")
    await state.set_state(AdminStates.wait_for_file_upload)

@router.message(AdminStates.wait_for_file_upload, F.document)
async def process_admin_file(message: Message, state: FSMContext, bot: Bot):
    doc = message.document
    status_msg = await message.reply("⏳ **ဖိုင်ကို ဖတ်ရှု၍ Gemini 2.5 Flash ဖြင့် မေးခွန်းထုတ်ပေးနေပါသည်...**")
    
    file_bytes = await bot.download(doc)
    text = extract_text_from_bytes(file_bytes.read(), doc.file_name)
    
    if not text.strip():
        await status_msg.edit_text("❌ ဖိုင်ထဲမှ စာသားထုတ်ယူ၍ မရရှိပါ။")
        await state.clear()
        return

    try:
        quiz_data: QuizSetSchema = await generate_quiz_from_gemini(text, count=5)
        default_cat = "General"
        for q in quiz_data.questions:
            await db_query(
                "INSERT INTO QuestionBank (category_name, question, options, correct_index, explanation) VALUES (?, ?, ?, ?, ?)",
                (default_cat, q.question, json.dumps(q.options), q.correct_index, q.explanation), commit=True
            )
        await status_msg.edit_text(f"✅ ဖိုင်ထဲမှ မေးခွန်း **({len(quiz_data.questions)})** ခုအား အောင်မြင်စွာ ထုတ်ယူသိမ်းဆည်းပြီးပါပြီ။", reply_markup=get_admin_menu())
    except Exception as e:
        await status_msg.edit_text(f"❌ Gemini AI Error: {str(e)}", reply_markup=get_admin_menu())
    
    await state.clear()

# ==========================================
# 5. BOT EXECUTION SETUP
# ==========================================
async def main():
    bot = Bot(token=TELEGRAM_TOKEN)
    dp = Dispatcher(storage=MemoryStorage())
    dp.include_router(router)
    
    print("🚀 Enterprise Quiz Bot Core is Running...")
    await dp.start_polling(bot)

if __name__ == "__main__":
    asyncio.run(main())
